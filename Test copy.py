from django.conf import settings
from openpyxl import load_workbook
import pandas as pd
from datetime import datetime
import oracledb
# Define your required columns
REQUIRED_COLUMNS = ["ARTICLE_CODE", "SU", "UNIQ_CODE", "UNIQ_NAME", "FROM_DATE", "TO_DATE","FLYER_RSP", "REG_RSP", "UNIT_DN", "REMARKS", "FLYER_TYPE", "CREATED_BY", "APPLICABLE_LOCATIONS"]
DATE_COLUMNS = ["FROM_DATE", "TO_DATE"]
REQUIRED_VALUE_COLUMNS = ["ARTICLE_CODE", "SU", "UNIQ_NAME", "UNIQ_CODE", "FROM_DATE", "TO_DATE", "FLYER_RSP", "FLYER_TYPE", "CREATED_BY"]
COLS_TO_FILL =[ "SU", "UNIQ_NAME", "FROM_DATE", "TO_DATE",  "REMARKS", "FLYER_TYPE", "CREATED_BY"]
INVALID_STRINGS = {"n/a", "na", "null", "none", "-"}
date_format = "%d-%b-%y"
path = r"C:\Users\HP\Downloads\testinh\1 page ok.xlsx"
# path = "TestFlareExcel.xlsx"
all_sheets = pd.read_excel(path, sheet_name=None)
# all_sheets = {"Sheet1": all_sheets["Sheet1"]}  # For testing, only use "Sheet1"
all_sheets_list = list(all_sheets.items())


section = [{'sheetname': 'Sheet1', 'removable_rows': [9, 19, 26, 27, 34], 'is_last_item_only_row': [21, 22, 23, 30, 31]}]


section_start_end = {}


def connection():

    username = 'KHYPER'
    password = 'KHYPER'
    dsn = '192.168.2.171:1521/ZEDEYE'
    client_path = "C:\instantclient_19_5"
    print(username)
    print(password)
    print(dsn)
    print(client_path)
    try:

        oracledb.init_oracle_client(lib_dir=client_path)

        conn = oracledb.connect(user=username, password=password, dsn=dsn)

        return conn
    except oracledb.Error as e:

        raise
    except Exception as e:
        raise


START_ROW = 2
all_rows_to_upload = []
for sheet_name, df in all_sheets_list:

    END_ROW = len(df) + 1  

    breaks = next(
    (s["removable_rows"] for s in section if s["sheetname"] == sheet_name),
    []
)
    is_last_item_only_row = next(
    (item['is_last_item_only_row'] for item in section if item['sheetname'] == sheet_name),
    []
)

    sections = []
    current_start = START_ROW

    for b in breaks:
        section_end = b - 1
        if current_start <= section_end:
            sections.append((current_start, section_end))
        current_start = b + 1

    # final section (dynamic end)
    if current_start <= END_ROW:
        sections.append((current_start, END_ROW))

    for start, end in sections:
        cutted = df.iloc[start-2:end-1].copy()  # make a copy to safely modify

        locations = cutted["APPLICABLE_LOCATIONS"].dropna().astype(
            int).tolist()
        print("Locations:", locations)

        cutted = cutted.drop(columns=["APPLICABLE_LOCATIONS"])
        # keep rows NOT in the list
        mask = ~cutted.index.isin(is_last_item_only_row)
        filtered_cutted = cutted.loc[mask].copy()

        for col in COLS_TO_FILL:
            filtered_cutted[col] = filtered_cutted[col].ffill()

        for index, row in filtered_cutted.iterrows():
            for location in locations:
                row_to_upload = row.copy()
                row_to_upload["APPLICABLE_LOCATIONS"] = location
                row_to_upload["CREATED_BY"] = "PERSON1"
                row_to_upload["SHEET"] = sheet_name
                all_rows_to_upload.append(row_to_upload)

conn = connection()
cursor = conn.cursor()

import math
from datetime import datetime
import pandas as pd 

# Function to clean values for Oracle
def clean_value(val):
    if isinstance(val, pd.Timestamp):  # convert pandas Timestamp to datetime
        return val.to_pydatetime()
    elif isinstance(val, datetime):  # datetime is fine as-is
        return val
    elif isinstance(val, float) and math.isnan(val):  # convert NaN to None
        return None
    else:
        return val

if all_rows_to_upload:  # make sure this list is not empty
    columns = list(all_rows_to_upload[0].keys())
    placeholders = ", ".join([f":{i+1}" for i in range(len(columns))])
    
    sql = f"""
    INSERT INTO KWT_FLYER_ART_UPLOAD_WKLY
    ({', '.join(columns)})
    VALUES ({placeholders})
    """
    
    try:
        for row in all_rows_to_upload:
            values = [clean_value(row[col]) for col in columns]
            print("Values:", values)
            cursor.execute(sql, values)
        conn.commit()
        print("All rows inserted successfully!")
    except Exception as e:
        conn.rollback()
        print("Error inserting rows:", e)
    finally:
        cursor.close()
        conn.close()
